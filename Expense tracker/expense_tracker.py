import csv
import os
import calendar
import datetime
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from expense import Expense

# Paths (Excel file in current directory)
EXPENSE_FILE = "expenses.csv"
HISTORY_FILE = "expense_history.xlsx"
BUDGET = 5000

def main():
    print("📘 Expense Tracker")

    handle_monthly_reset()

    # Ensure the CSV file is initialized
    if not os.path.exists(EXPENSE_FILE):
        with open(EXPENSE_FILE, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["name", "amount", "category", "date"])

    while True:
        print("\nMenu:")
        print("1. Add New Expense")
        print("2. View Summary")
        print("3. Edit Expense")
        print("4. Delete Expense")
        print("5. Exit")
        choice = input("Select an option (1/2/3/4/5): ")

        if choice == "1":
            expense = get_user_expense()
            save_expense_to_csv(expense)
            save_expense_to_excel(expense)
        elif choice == "2":
            summarize_expenses(EXPENSE_FILE, BUDGET)
        elif choice == "3":
            edit_expense(EXPENSE_FILE)
        elif choice == "4":
            delete_expense(EXPENSE_FILE)
        elif choice == "5":
            print("👋 Exiting Expense Tracker. Goodbye!")
            break
        else:
            print("❌ Invalid option. Please choose again.")

def handle_monthly_reset():
    if not os.path.exists(EXPENSE_FILE):
        return

    with open(EXPENSE_FILE, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        data = list(reader)

    if not data:
        return

    last_date = data[-1]["date"]
    last_datetime = datetime.datetime.strptime(last_date, "%Y-%m-%d")
    last_month = last_datetime.month
    last_year = last_datetime.year

    current_month = datetime.datetime.now().month
    current_year = datetime.datetime.now().year

    if last_month != current_month or last_year != current_year:
        sheet_name = f"{calendar.month_name[last_month]}_{last_year}"
        save_to_excel(data, sheet_name=sheet_name)

        with open(EXPENSE_FILE, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["name", "amount", "category", "date"])
        print(f"📁 Data for {sheet_name} moved to '{HISTORY_FILE}' and new file started.")

def save_to_excel(data, sheet_name):
    if os.path.exists(HISTORY_FILE):
        wb = load_workbook(HISTORY_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name in wb.sheetnames:
        base_name = sheet_name
        counter = 1
        while sheet_name in wb.sheetnames:
            sheet_name = f"{base_name}_{counter}"
            counter += 1

    ws = wb.create_sheet(title=sheet_name)
    ws.append(["name", "amount", "category", "date"])
    for row in data:
        ws.append([row["name"], float(row["amount"]), row["category"], row["date"]])

    wb.save(HISTORY_FILE)

def save_expense_to_csv(expense):
    with open(EXPENSE_FILE, "a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([expense.name, expense.amount, expense.category, expense.date])

def save_expense_to_excel(expense):
    month_name = calendar.month_name[datetime.datetime.now().month]
    year = datetime.datetime.now().year
    sheet_name = f"{month_name}_{year}"

    if os.path.exists(HISTORY_FILE):
        wb = load_workbook(HISTORY_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["name", "amount", "category", "date"])
    else:
        ws = wb[sheet_name]

    ws.append([expense.name, expense.amount, expense.category, expense.date])
    wb.save(HISTORY_FILE)

    print(f"✅ Saved to Excel under sheet: {sheet_name}")

def get_user_expense():
    print("\n📝 Entering New Expense")
    name = input("Enter expense name: ")
    while True:
        try:
            amount = float(input("Enter expense amount: "))
            break
        except ValueError:
            print("❌ Invalid amount. Try again.")

    categories = ["Food", "Milk & Snacks", "Bus", "Grocery", "Entertainment", "Stationary", "Clothing", "Other"]
    for i, category in enumerate(categories, start=1):
        print(f"{i}. {category}")

    while True:
        try:
            idx = int(input(f"Select category (1-{len(categories)}): ")) - 1
            if 0 <= idx < len(categories):
                category = categories[idx]
                today = datetime.date.today().isoformat()
                return Expense(name=name, amount=amount, category=category, date=today)
            else:
                print("❌ Invalid category.")
        except ValueError:
            print("❌ Enter a number.")

def summarize_expenses(path, budget):
    if not os.path.exists(path):
        print("⚠️ No expenses recorded.")
        return

    with open(path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        expenses = [Expense(row["name"], float(row["amount"]), row["category"], row["date"]) for row in reader]

    if not expenses:
        print("⚠️ No valid expenses found.")
        return

    print("\n📊 Expense Summary")
    total_spent = sum(e.amount for e in expenses)
    remaining = budget - total_spent
    percent_used = (total_spent / budget) * 100

    category_totals = {}
    for e in expenses:
        category_totals[e.category] = category_totals.get(e.category, 0) + e.amount

    for cat, amt in category_totals.items():
        print(f"  {cat}: ₹{amt:.2f}")

    print(f"\nTotal Spent: ₹{total_spent:.2f}")
    print(f"Remaining Budget: ₹{remaining:.2f}")
    print(f"Budget Used: {percent_used:.2f}%")

    now = datetime.datetime.now()
    days_left = calendar.monthrange(now.year, now.month)[1] - now.day
    print(f"Daily Budget: ₹{remaining / max(days_left, 1):.2f}")

    category_totals["Remaining"] = remaining
    plot_expenses(category_totals, total_spent, remaining)

def plot_expenses(category_totals, total_spent, remaining):
    print("📈 Generating graphs...")
    categories = list(category_totals.keys())
    values = list(category_totals.values())

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 6))
    ax1.bar(categories, values, color='skyblue')
    ax1.set_title("Expenses + Remaining Budget")
    ax1.set_xlabel("Category")
    ax1.set_ylabel("Amount (₹)")
    ax1.tick_params(axis='x', rotation=45)

    ax2.pie(values, labels=categories, autopct='%1.1f%%', startangle=140)
    ax2.set_title("Pie Chart with Remaining Budget")
    ax2.text(0.5, -1.3, f"Total Spent: ₹{total_spent:.2f}\nRemaining: ₹{remaining:.2f}", fontsize=12, ha='center')

    fig.tight_layout()
    plt.show()

def delete_expense(path):
    if not os.path.exists(path):
        print("⚠️ Expense file not found.")
        return

    with open(path, "r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))

    if not reader:
        print("⚠️ No expenses to delete.")
        return

    print("\n🗑️ Select expense to delete:")
    for i, row in enumerate(reader):
        print(f"{i+1}. {row['name']} | ₹{row['amount']} | {row['category']} | {row['date']}")

    try:
        idx = int(input("Enter the number to delete: ")) - 1
        if 0 <= idx < len(reader):
            deleted = reader.pop(idx)
            print(f"✅ Deleted: {deleted['name']} | ₹{deleted['amount']}")
        else:
            print("❌ Invalid index.")
            return
    except ValueError:
        print("❌ Enter a valid number.")
        return

    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "amount", "category", "date"])
        writer.writeheader()
        writer.writerows(reader)

def edit_expense(path):
    if not os.path.exists(path):
        print("⚠️ Expense file not found.")
        return

    with open(path, "r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))

    if not reader:
        print("⚠️ No expenses to edit.")
        return

    print("\n✏️ Select expense to edit:")
    for i, row in enumerate(reader):
        print(f"{i+1}. {row['name']} | ₹{row['amount']} | {row['category']} | {row['date']}")

    try:
        idx = int(input("Enter the number to edit: ")) - 1
        if 0 <= idx < len(reader):
            row = reader[idx]
            print("Leave blank to keep current value.")
            new_name = input(f"Name ({row['name']}): ") or row['name']
            new_amount = input(f"Amount ({row['amount']}): ") or row['amount']
            new_category = input(f"Category ({row['category']}): ") or row['category']
            new_date = input(f"Date ({row['date']}): ") or row['date']

            reader[idx] = {
                "name": new_name,
                "amount": new_amount,
                "category": new_category,
                "date": new_date
            }

            print("✅ Expense updated.")
        else:
            print("❌ Invalid index.")
            return
    except ValueError:
        print("❌ Enter a valid number.")
        return

    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "amount", "category", "date"])
        writer.writeheader()
        writer.writerows(reader)

if __name__ == "__main__":
    main()
# This code is a simple expense tracker that allows users to add, view, edit, and delete expenses.
# It saves expenses to a CSV file and also maintains an Excel history of monthly expenses.          
# It provides a summary of expenses, including total spent, remaining budget, and daily budget.
# It also generates bar and pie charts to visualize expenses by category.       \
