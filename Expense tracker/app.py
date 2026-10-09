import streamlit as st
import pandas as pd
import csv
import os
import calendar
import datetime
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from expense import Expense

# Paths
EXPENSE_FILE = "expenses.csv"
HISTORY_FILE = "expense_history.xlsx"
BUDGET = 5000

st.set_page_config(page_title="Expense Tracker", page_icon="💸", layout="centered")

def handle_monthly_reset():
    if not os.path.exists(EXPENSE_FILE):
        return

    with open(EXPENSE_FILE, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        data = list(reader)

    if not data:
        return

    last_date = data[-1]["date"]
    last_datetime = datetime.datetime.strptime(last_date, "%Y-%m-%d")
    last_month = last_datetime.month
    last_year = last_datetime.year

    current_month = datetime.datetime.now().month
    current_year = datetime.datetime.now().year

    if last_month != current_month or last_year != current_year:
        sheet_name = f"{calendar.month_name[last_month]}_{last_year}"
        save_to_excel(data, sheet_name=sheet_name)

        with open(EXPENSE_FILE, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["name", "amount", "category", "date"])
        st.info(f"Data for {sheet_name} moved to '{HISTORY_FILE}' and new file started.")

def save_to_excel(data, sheet_name):
    if os.path.exists(HISTORY_FILE):
        wb = load_workbook(HISTORY_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name in wb.sheetnames:
        base_name = sheet_name
        counter = 1
        while sheet_name in wb.sheetnames:
            sheet_name = f"{base_name}_{counter}"
            counter += 1

    ws = wb.create_sheet(title=sheet_name)
    ws.append(["name", "amount", "category", "date"])
    for row in data:
        ws.append([row["name"], float(row["amount"]), row["category"], row["date"]])

    wb.save(HISTORY_FILE)

def save_expense_to_csv(expense):
    with open(EXPENSE_FILE, "a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([expense.name, expense.amount, expense.category, expense.date])

def save_expense_to_excel(expense):
    month_name = calendar.month_name[datetime.datetime.now().month]
    year = datetime.datetime.now().year
    sheet_name = f"{month_name}_{year}"

    if os.path.exists(HISTORY_FILE):
        wb = load_workbook(HISTORY_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["name", "amount", "category", "date"])
    else:
        ws = wb[sheet_name]

    ws.append([expense.name, expense.amount, expense.category, expense.date])
    wb.save(HISTORY_FILE)

# Ensure the CSV file is initialized
if not os.path.exists(EXPENSE_FILE):
    with open(EXPENSE_FILE, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["name", "amount", "category", "date"])

handle_monthly_reset()

st.title("💸 Expense Tracker")

menu = ["Add New Expense", "View Summary", "Edit Expense", "Delete Expense"]
choice = st.sidebar.selectbox("Menu", menu)

categories = ["Food", "Milk & Snacks", "Bus", "Grocery", "Entertainment", "Stationary", "Clothing", "Other"]

if choice == "Add New Expense":
    st.header("Add New Expense")
    with st.form("expense_form"):
        name = st.text_input("Expense Name")
        amount = st.number_input("Amount (₹)", min_value=0.0, format="%.2f")
        category = st.selectbox("Category", categories)
        date = st.date_input("Date")
        submitted = st.form_submit_button("Add Expense")
        
        if submitted:
            if name and amount > 0:
                expense = Expense(name=name, amount=amount, category=category, date=date.isoformat())
                save_expense_to_csv(expense)
                save_expense_to_excel(expense)
                st.success(f"Expense '{name}' added successfully!")
            else:
                st.error("Please enter a valid name and amount.")

elif choice == "View Summary":
    st.header("Expense Summary")
    if os.path.exists(EXPENSE_FILE):
        df = pd.read_csv(EXPENSE_FILE)
        if not df.empty:
            st.dataframe(df)
            
            total_spent = df["amount"].sum()
            remaining = BUDGET - total_spent
            percent_used = (total_spent / BUDGET) * 100
            
            st.write(f"**Total Spent:** ₹{total_spent:.2f}")
            st.write(f"**Remaining Budget:** ₹{remaining:.2f}")
            st.write(f"**Budget Used:** {percent_used:.2f}%")
            
            now = datetime.datetime.now()
            days_left = calendar.monthrange(now.year, now.month)[1] - now.day
            st.write(f"**Daily Budget for remaining days:** ₹{remaining / max(days_left, 1):.2f}")
            
            category_totals = df.groupby("category")["amount"].sum().to_dict()
            category_totals["Remaining"] = max(remaining, 0)
            
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 6))
            cats = [c for c in category_totals.keys() if c != "Remaining"]
            vals = [category_totals[c] for c in cats]
            
            ax1.bar(cats, vals, color='skyblue')
            ax1.set_title("Expenses by Category")
            ax1.set_xlabel("Category")
            ax1.set_ylabel("Amount (₹)")
            ax1.tick_params(axis='x', rotation=45)
            
            ax2.pie(category_totals.values(), labels=category_totals.keys(), autopct='%1.1f%%', startangle=140)
            ax2.set_title("Budget Breakdown")
            
            st.pyplot(fig)
        else:
            st.info("No expenses recorded yet.")
    else:
        st.info("No expenses recorded yet.")

elif choice == "Edit Expense":
    st.header("Edit Expense")
    if os.path.exists(EXPENSE_FILE):
        df = pd.read_csv(EXPENSE_FILE)
        if not df.empty:
            expense_to_edit = st.selectbox("Select expense to edit", df.index, format_func=lambda i: f"{df.loc[i, 'name']} - ₹{df.loc[i, 'amount']} ({df.loc[i, 'date']})")
            
            with st.form("edit_form"):
                row = df.loc[expense_to_edit]
                new_name = st.text_input("Name", value=row["name"])
                new_amount = st.number_input("Amount (₹)", value=float(row["amount"]), min_value=0.0, format="%.2f")
                new_category = st.selectbox("Category", categories, index=categories.index(row["category"]) if row["category"] in categories else 0)
                try:
                    curr_date = datetime.datetime.strptime(row["date"], "%Y-%m-%d").date()
                except:
                    curr_date = datetime.date.today()
                new_date = st.date_input("Date", value=curr_date)
                
                submitted = st.form_submit_button("Update Expense")
                if submitted:
                    df.at[expense_to_edit, "name"] = new_name
                    df.at[expense_to_edit, "amount"] = new_amount
                    df.at[expense_to_edit, "category"] = new_category
                    df.at[expense_to_edit, "date"] = new_date.isoformat()
                    df.to_csv(EXPENSE_FILE, index=False)
                    st.success("Expense updated successfully!")
        else:
            st.info("No expenses to edit.")
    else:
        st.info("No expenses to edit.")

elif choice == "Delete Expense":
    st.header("Delete Expense")
    if os.path.exists(EXPENSE_FILE):
        df = pd.read_csv(EXPENSE_FILE)
        if not df.empty:
            expense_to_delete = st.selectbox("Select expense to delete", df.index, format_func=lambda i: f"{df.loc[i, 'name']} - ₹{df.loc[i, 'amount']} ({df.loc[i, 'date']})")
            
            if st.button("Delete Expense"):
                df = df.drop(index=expense_to_delete)
                df.to_csv(EXPENSE_FILE, index=False)
                st.success("Expense deleted successfully!")
                st.rerun()
        else:
            st.info("No expenses to delete.")
    else:
        st.info("No expenses to delete.")
